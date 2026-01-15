import json
from dataclasses import dataclass, field
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook

CONFIG_PATH = Path(__file__).resolve().parent.parent / "config" / "templates.json"

AI_HEADERS = [
    "№ п/п",
    "Наименование присоединения",
    "Наименование сигнала",
    "Тип сигнала",
    "Ед. измер.",
    "Уровень сигнала",
    "Место съема сигнала",
    "Устройство регистрации сигнала",
]

DI_HEADERS = [
    "№ п/п",
    "Наименование присоединения",
    "Наименование аппарата",
    "Наименование сигнала",
    "Тип сигнала",
    "Хар-ка сигнала",
    "Место съема сигнала",
    "Устройство регистрации сигнала",
]

DO_HEADERS = [
    "№ п/п",
    "Наименование присоединения",
    "Наименование аппарата",
    "Наименование сигнала",
    "Тип сигнала",
    "Хар-ка сигнала",
    "Место получения сигнала",
    "Устройство формирования сигнала",
]


@dataclass
class ConnectionRow:
    name_var: tk.StringVar
    type_var: tk.StringVar
    include_ai: tk.BooleanVar
    include_di: tk.BooleanVar
    include_do: tk.BooleanVar


@dataclass
class VoltageClassSection:
    name_var: tk.StringVar
    rows: list[ConnectionRow] = field(default_factory=list)


class SignalGeneratorApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Генератор перечня сигналов АСУ ТП")
        self.geometry("1100x700")

        self.templates = self._load_templates()
        self.type_names = sorted(self.templates["types"].keys())
        self.sections: list[VoltageClassSection] = []

        self._build_ui()

    def _load_templates(self) -> dict:
        if not CONFIG_PATH.exists():
            raise FileNotFoundError(f"Не найден файл конфигурации: {CONFIG_PATH}")
        with CONFIG_PATH.open("r", encoding="utf-8") as handle:
            return json.load(handle)

    def _build_ui(self) -> None:
        control_frame = ttk.Frame(self)
        control_frame.pack(fill=tk.X, padx=12, pady=10)

        ttk.Label(control_frame, text="Класс напряжения").grid(row=0, column=0, sticky=tk.W)
        self.voltage_name_var = tk.StringVar()
        ttk.Entry(control_frame, textvariable=self.voltage_name_var, width=20).grid(
            row=0, column=1, padx=6
        )

        ttk.Label(control_frame, text="Количество присоединений").grid(row=0, column=2, sticky=tk.W)
        self.voltage_count_var = tk.StringVar(value="1")
        ttk.Entry(control_frame, textvariable=self.voltage_count_var, width=6).grid(
            row=0, column=3, padx=6
        )

        ttk.Button(
            control_frame,
            text="Добавить класс",
            command=self._add_voltage_class,
        ).grid(row=0, column=4, padx=6)

        ttk.Button(
            control_frame,
            text="Сформировать Excel",
            command=self._generate_excel,
        ).grid(row=0, column=5, padx=6)

        ttk.Button(
            control_frame,
            text="Сохранить проект",
            command=self._save_project,
        ).grid(row=0, column=6, padx=6)

        ttk.Button(
            control_frame,
            text="Загрузить проект",
            command=self._load_project,
        ).grid(row=0, column=7, padx=6)

        ttk.Button(
            control_frame,
            text="Очистить",
            command=self._clear_form,
        ).grid(row=0, column=8, padx=6)

        self.canvas = tk.Canvas(self, borderwidth=0)
        self.scroll_frame = ttk.Frame(self.canvas)
        self.scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.scrollbar.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        self.scroll_frame.bind("<Configure>", self._on_frame_configure)

    def _on_frame_configure(self, _event: tk.Event) -> None:
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _add_voltage_class(self) -> None:
        name = self.voltage_name_var.get().strip()
        count_str = self.voltage_count_var.get().strip()
        if not name:
            messagebox.showerror("Ошибка", "Укажите класс напряжения.")
            return
        try:
            count = int(count_str)
        except ValueError:
            messagebox.showerror("Ошибка", "Количество присоединений должно быть числом.")
            return
        if count <= 0:
            messagebox.showerror("Ошибка", "Количество присоединений должно быть больше 0.")
            return

        section = VoltageClassSection(name_var=tk.StringVar(value=name))
        self.sections.append(section)
        section_frame = ttk.LabelFrame(self.scroll_frame, text=f"Класс {name}")
        section_frame.pack(fill=tk.X, padx=12, pady=8, anchor="n")

        headers = [
            "Наименование присоединения",
            "Тип присоединения",
            "AI",
            "DI",
            "DO",
        ]
        for col_index, header in enumerate(headers):
            ttk.Label(section_frame, text=header).grid(row=0, column=col_index, padx=6, pady=4)

        for row_index in range(count):
            name_var = tk.StringVar()
            type_var = tk.StringVar()
            include_ai = tk.BooleanVar(value=True)
            include_di = tk.BooleanVar(value=True)
            include_do = tk.BooleanVar(value=True)

            ttk.Entry(section_frame, textvariable=name_var, width=40).grid(
                row=row_index + 1, column=0, padx=6, pady=2, sticky=tk.W
            )
            type_combo = ttk.Combobox(
                section_frame,
                textvariable=type_var,
                values=self.type_names,
                state="readonly",
                width=25,
            )
            type_combo.grid(row=row_index + 1, column=1, padx=6, pady=2)

            ttk.Checkbutton(section_frame, variable=include_ai).grid(
                row=row_index + 1, column=2
            )
            ttk.Checkbutton(section_frame, variable=include_di).grid(
                row=row_index + 1, column=3
            )
            ttk.Checkbutton(section_frame, variable=include_do).grid(
                row=row_index + 1, column=4
            )

            section.rows.append(
                ConnectionRow(
                    name_var=name_var,
                    type_var=type_var,
                    include_ai=include_ai,
                    include_di=include_di,
                    include_do=include_do,
                )
            )

        self.voltage_name_var.set("")
        self.voltage_count_var.set("1")

    def _collect_data(self) -> tuple[list[dict], list[str]]:
        data = []
        errors = []
        template_types = self.templates["types"]

        for section_index, section in enumerate(self.sections, start=1):
            class_name = section.name_var.get().strip()
            if not class_name:
                errors.append(f"Класс напряжения #{section_index} не заполнен.")
                continue

            for row_index, row in enumerate(section.rows, start=1):
                conn_name = row.name_var.get().strip()
                conn_type = row.type_var.get().strip()
                if not conn_name:
                    errors.append(
                        f"Пустое наименование присоединения: {class_name}, строка {row_index}."
                    )
                if not conn_type:
                    errors.append(
                        f"Не выбран тип присоединения: {class_name}, строка {row_index}."
                    )
                elif conn_type not in template_types:
                    errors.append(
                        f"Нет шаблона для типа '{conn_type}': {class_name}, строка {row_index}."
                    )

                data.append(
                    {
                        "class": class_name,
                        "name": conn_name,
                        "type": conn_type,
                        "include_ai": row.include_ai.get(),
                        "include_di": row.include_di.get(),
                        "include_do": row.include_do.get(),
                    }
                )
        return data, errors

    def _apply_defaults(self, signal: dict, defaults: dict) -> dict:
        filled = defaults.copy()
        filled.update(signal)
        return filled

    def _generate_excel(self) -> None:
        data, errors = self._collect_data()
        if errors:
            messagebox.showerror("Ошибки проверки", "\n".join(errors))
            return

        save_path = filedialog.asksaveasfilename(
            title="Сохранить Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not save_path:
            return

        workbook = Workbook()
        ai_sheet = workbook.active
        ai_sheet.title = "AI"
        di_sheet = workbook.create_sheet("DI")
        do_sheet = workbook.create_sheet("DO")

        ai_sheet.append(AI_HEADERS)
        di_sheet.append(DI_HEADERS)
        do_sheet.append(DO_HEADERS)

        ai_index = 1
        di_index = 1
        do_index = 1

        defaults = self.templates.get("defaults", {})

        for item in data:
            template = self.templates["types"].get(item["type"], {})
            if item["include_ai"]:
                for signal in template.get("ai", []):
                    values = self._apply_defaults(signal, defaults.get("ai", {}))
                    ai_sheet.append(
                        [
                            ai_index,
                            item["name"],
                            values.get("Наименование сигнала", ""),
                            values.get("Тип сигнала", ""),
                            values.get("Ед. измер.", ""),
                            values.get("Уровень сигнала", ""),
                            values.get("Место съема сигнала", ""),
                            values.get("Устройство регистрации сигнала", ""),
                        ]
                    )
                    ai_index += 1

            if item["include_di"]:
                for signal in template.get("di", []):
                    values = self._apply_defaults(signal, defaults.get("di", {}))
                    di_sheet.append(
                        [
                            di_index,
                            item["name"],
                            values.get("Наименование аппарата", ""),
                            values.get("Наименование сигнала", ""),
                            values.get("Тип сигнала", ""),
                            values.get("Хар-ка сигнала", ""),
                            values.get("Место съема сигнала", ""),
                            values.get("Устройство регистрации сигнала", ""),
                        ]
                    )
                    di_index += 1

            if item["include_do"]:
                for signal in template.get("do", []):
                    values = self._apply_defaults(signal, defaults.get("do", {}))
                    do_sheet.append(
                        [
                            do_index,
                            item["name"],
                            values.get("Наименование аппарата", ""),
                            values.get("Наименование сигнала", ""),
                            values.get("Тип сигнала", ""),
                            values.get("Хар-ка сигнала", ""),
                            values.get("Место получения сигнала", ""),
                            values.get("Устройство формирования сигнала", ""),
                        ]
                    )
                    do_index += 1

        workbook.save(save_path)
        messagebox.showinfo("Готово", f"Файл сохранен: {save_path}")

    def _clear_form(self) -> None:
        for widget in self.scroll_frame.winfo_children():
            widget.destroy()
        self.sections.clear()

    def _save_project(self) -> None:
        data, errors = self._collect_data()
        if errors:
            messagebox.showerror("Ошибки проверки", "\n".join(errors))
            return

        save_path = filedialog.asksaveasfilename(
            title="Сохранить проект",
            defaultextension=".json",
            filetypes=[("JSON", "*.json")],
        )
        if not save_path:
            return

        project = {"classes": []}
        for section in self.sections:
            class_data = {
                "name": section.name_var.get().strip(),
                "connections": [],
            }
            for row in section.rows:
                class_data["connections"].append(
                    {
                        "name": row.name_var.get().strip(),
                        "type": row.type_var.get().strip(),
                        "include_ai": row.include_ai.get(),
                        "include_di": row.include_di.get(),
                        "include_do": row.include_do.get(),
                    }
                )
            project["classes"].append(class_data)

        with open(save_path, "w", encoding="utf-8") as handle:
            json.dump(project, handle, ensure_ascii=False, indent=2)

        messagebox.showinfo("Готово", f"Проект сохранен: {save_path}")

    def _load_project(self) -> None:
        open_path = filedialog.askopenfilename(
            title="Загрузить проект",
            filetypes=[("JSON", "*.json")],
        )
        if not open_path:
            return

        try:
            with open(open_path, "r", encoding="utf-8") as handle:
                project = json.load(handle)
        except (OSError, json.JSONDecodeError) as exc:
            messagebox.showerror("Ошибка", f"Не удалось прочитать файл: {exc}")
            return

        self._clear_form()

        for class_data in project.get("classes", []):
            name = class_data.get("name", "")
            connections = class_data.get("connections", [])
            section = VoltageClassSection(name_var=tk.StringVar(value=name))
            self.sections.append(section)

            section_frame = ttk.LabelFrame(self.scroll_frame, text=f"Класс {name}")
            section_frame.pack(fill=tk.X, padx=12, pady=8, anchor="n")

            headers = [
                "Наименование присоединения",
                "Тип присоединения",
                "AI",
                "DI",
                "DO",
            ]
            for col_index, header in enumerate(headers):
                ttk.Label(section_frame, text=header).grid(
                    row=0, column=col_index, padx=6, pady=4
                )

            for row_index, connection in enumerate(connections, start=1):
                name_var = tk.StringVar(value=connection.get("name", ""))
                type_var = tk.StringVar(value=connection.get("type", ""))
                include_ai = tk.BooleanVar(value=connection.get("include_ai", True))
                include_di = tk.BooleanVar(value=connection.get("include_di", True))
                include_do = tk.BooleanVar(value=connection.get("include_do", True))

                ttk.Entry(section_frame, textvariable=name_var, width=40).grid(
                    row=row_index, column=0, padx=6, pady=2, sticky=tk.W
                )
                type_combo = ttk.Combobox(
                    section_frame,
                    textvariable=type_var,
                    values=self.type_names,
                    state="readonly",
                    width=25,
                )
                type_combo.grid(row=row_index, column=1, padx=6, pady=2)

                ttk.Checkbutton(section_frame, variable=include_ai).grid(
                    row=row_index, column=2
                )
                ttk.Checkbutton(section_frame, variable=include_di).grid(
                    row=row_index, column=3
                )
                ttk.Checkbutton(section_frame, variable=include_do).grid(
                    row=row_index, column=4
                )

                section.rows.append(
                    ConnectionRow(
                        name_var=name_var,
                        type_var=type_var,
                        include_ai=include_ai,
                        include_di=include_di,
                        include_do=include_do,
                    )
                )

    def run(self) -> None:
        self.mainloop()


if __name__ == "__main__":
    app = SignalGeneratorApp()
    app.run()
